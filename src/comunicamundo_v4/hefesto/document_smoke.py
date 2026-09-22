from __future__ import annotations

import json
from pathlib import Path
import subprocess
from tempfile import TemporaryDirectory
from zipfile import ZIP_DEFLATED, ZIP_STORED, ZipFile

from PIL import Image, ImageDraw

from .document_formats import resolve_document_format
from .document_normalization import _document_converter, _status_name


def _write_office_files(root: Path) -> list[Path]:
    from docx import Document
    from openpyxl import Workbook
    from pptx import Presentation

    docx = root / "source.docx"
    document = Document()
    document.add_heading("Documento de fumaça", level=1)
    document.add_paragraph("Um parágrafo verificável para a entrada Word.")
    document.save(docx)

    pptx = root / "source.pptx"
    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[1])
    slide.shapes.title.text = "Apresentação de fumaça"
    slide.placeholders[1].text = "Conteúdo verificável no PowerPoint."
    presentation.save(pptx)

    xlsx = root / "source.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.append(["Assunto", "Valor"])
    sheet.append(["Entrada Excel", 1])
    workbook.save(xlsx)
    return [docx, pptx, xlsx]


def _write_epub(path: Path) -> None:
    with ZipFile(path, "w") as archive:
        archive.writestr("mimetype", "application/epub+zip", compress_type=ZIP_STORED)
        archive.writestr(
            "META-INF/container.xml",
            """<?xml version="1.0"?><container version="1.0" xmlns="urn:oasis:names:tc:opendocument:xmlns:container"><rootfiles><rootfile full-path="OEBPS/content.opf" media-type="application/oebps-package+xml"/></rootfiles></container>""",
            compress_type=ZIP_DEFLATED,
        )
        archive.writestr(
            "OEBPS/content.opf",
            """<?xml version="1.0"?><package version="3.0" xmlns="http://www.idpf.org/2007/opf" unique-identifier="id"><metadata xmlns:dc="http://purl.org/dc/elements/1.1/"><dc:identifier id="id">smoke</dc:identifier><dc:title>EPUB smoke</dc:title><dc:language>pt-BR</dc:language></metadata><manifest><item id="chapter" href="chapter.xhtml" media-type="application/xhtml+xml"/></manifest><spine><itemref idref="chapter"/></spine></package>""",
            compress_type=ZIP_DEFLATED,
        )
        archive.writestr(
            "OEBPS/chapter.xhtml",
            """<html xmlns="http://www.w3.org/1999/xhtml"><head><title>EPUB</title></head><body><h1>Livro de fumaça</h1><p>Conteúdo verificável no EPUB.</p></body></html>""",
            compress_type=ZIP_DEFLATED,
        )


def _write_media(root: Path) -> list[Path]:
    wav = root / "source.wav"
    video = root / "source.mp4"
    subprocess.run(
        ["ffmpeg", "-loglevel", "error", "-f", "lavfi", "-i", "sine=frequency=440:duration=0.25", str(wav)],
        check=True,
    )
    subprocess.run(
        [
            "ffmpeg", "-loglevel", "error", "-f", "lavfi", "-i",
            "color=c=blue:s=320x240:d=0.25", "-f", "lavfi", "-i",
            "sine=frequency=440:duration=0.25", "-shortest", "-pix_fmt", "yuv420p",
            str(video),
        ],
        check=True,
    )
    return [wav, video]


def _write_pdf(path: Path) -> None:
    content = b"BT /F1 12 Tf 30 250 Td (PDF smoke content.) Tj ET"
    objects = [
        b"<< /Type /Catalog /Pages 2 0 R >>",
        b"<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
        b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 300 300] /Resources << /Font << /F1 4 0 R >> >> /Contents 5 0 R >>",
        b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>",
        b"<< /Length " + str(len(content)).encode() + b" >>\nstream\n" + content + b"\nendstream",
    ]
    data = bytearray(b"%PDF-1.4\n")
    offsets: list[int] = []
    for ordinal, value in enumerate(objects, start=1):
        offsets.append(len(data))
        data.extend(f"{ordinal} 0 obj\n".encode() + value + b"\nendobj\n")
    xref = len(data)
    data.extend(f"xref\n0 {len(objects) + 1}\n".encode())
    data.extend(b"0000000000 65535 f \n")
    for offset in offsets:
        data.extend(f"{offset:010d} 00000 n \n".encode())
    data.extend(f"trailer\n<< /Size 6 /Root 1 0 R >>\nstartxref\n{xref}\n%%EOF\n".encode())
    path.write_bytes(data)


def _fixtures(root: Path) -> list[Path]:
    pdf = root / "source.pdf"
    _write_pdf(pdf)
    html = root / "source.html"
    html.write_text("<h1>HTML de fumaça</h1><p>Conteúdo verificável.</p>", encoding="utf-8")
    markdown = root / "source.md"
    markdown.write_text("# Markdown de fumaça\n\nConteúdo verificável.", encoding="utf-8")
    text = root / "source.txt"
    text.write_text("Texto simples verificável.", encoding="utf-8")
    epub = root / "source.epub"
    _write_epub(epub)
    image = root / "source.png"
    canvas = Image.new("RGB", (900, 240), "white")
    ImageDraw.Draw(canvas).text((40, 90), "IMAGEM DE FUMACA COM TEXTO", fill="black")
    canvas.save(image)
    return [pdf, html, markdown, text, epub, image, *_write_office_files(root), *_write_media(root)]


def main() -> None:
    results: dict[str, str] = {}
    with TemporaryDirectory() as temporary:
        paths = _fixtures(Path(temporary))
        converter = _document_converter()
        for path in paths:
            document_format = resolve_document_format(path.name)
            if document_format.family == "text":
                if not path.read_text(encoding="utf-8").strip():
                    raise RuntimeError("plain-text smoke produced no text")
                results[path.suffix] = "success"
                continue
            converted = converter.convert(path, raises_on_error=False)
            status = _status_name(converted.status)
            if status not in {"success", "partial_success"}:
                raise RuntimeError(f"{path.suffix} smoke failed: {converted.errors}")
            results[path.suffix] = status
    print(json.dumps(results, ensure_ascii=False, sort_keys=True))


if __name__ == "__main__":
    main()
